import io
import os
import pytest
from PIL import Image
from docx import Document
from app import create_app
from config import TestingConfig


@pytest.fixture
def client():
    app = create_app(TestingConfig)
    with app.test_client() as client:
        yield client


def test_convert_missing_files(client):
    """Test POST /api/convert without files returns 400."""
    res = client.post("/api/convert", data={"conversion_type": "images_to_pdf"})
    assert res.status_code == 400
    data = res.get_json()
    assert data["success"] is False
    assert data["error"]["code"] == "MISSING_FILES"


def test_convert_missing_type(client):
    """Test POST /api/convert without conversion_type returns 400."""
    data = {
        "files": (io.BytesIO(b"dummy data"), "test.txt")
    }
    res = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert res.status_code == 400
    res_data = res.get_json()
    assert res_data["success"] is False
    assert res_data["error"]["code"] == "MISSING_CONVERSION_TYPE"


def test_convert_invalid_type(client):
    """Test POST /api/convert with unsupported conversion_type returns 422."""
    data = {
        "conversion_type": "unknown_type",
        "files": (io.BytesIO(b"dummy data"), "test.txt")
    }
    res = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert res.status_code == 422
    res_data = res.get_json()
    assert res_data["success"] is False
    assert res_data["error"]["code"] == "CONVERSION_FAILED"


def test_convert_images_to_pdf_endpoint(client):
    """Test uploading 2 images and converting to PDF."""
    # Create two 100x100 dummy images
    img1_io = io.BytesIO()
    Image.new("RGB", (100, 100), color="blue").save(img1_io, format="PNG")
    img1_io.seek(0)

    img2_io = io.BytesIO()
    Image.new("RGB", (100, 100), color="red").save(img2_io, format="JPEG")
    img2_io.seek(0)

    data = {
        "conversion_type": "images_to_pdf",
        "output_format": "pdf",
        "custom_filename": "test_portfolio",
        "files": [
            (img1_io, "page1.png"),
            (img2_io, "page2.jpg")
        ]
    }

    res = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert res.status_code == 200
    res_json = res.get_json()
    assert res_json["success"] is True
    assert res_json["filename"] == "test_portfolio.pdf"
    assert "file_id" in res_json

    # Test downloading the created file
    dl_res = client.get(f"/download/{res_json['file_id']}")
    assert dl_res.status_code == 200
    assert dl_res.data.startswith(b"%PDF")


def test_convert_pdf_to_images_endpoint(client):
    """Test uploading a PDF and converting to JPG images."""
    import fitz

    # Generate a simple 2-page PDF in memory
    doc = fitz.open()
    page1 = doc.new_page(width=300, height=300)
    page1.insert_text((50, 50), "Page 1 Content", fontsize=16)
    page2 = doc.new_page(width=300, height=300)
    page2.insert_text((50, 50), "Page 2 Content", fontsize=16)

    pdf_bytes = doc.write()
    doc.close()

    data = {
        "conversion_type": "pdf_to_images",
        "output_format": "jpg",
        "custom_filename": "my_pdf_pages",
        "files": (io.BytesIO(pdf_bytes), "test_doc.pdf")
    }

    res = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert res.status_code == 200
    res_json = res.get_json()
    assert res_json["success"] is True
    # For multi-page PDF, should be a ZIP of images
    assert res_json["filename"].endswith(".zip")
    assert "file_id" in res_json

    dl_res = client.get(f"/download/{res_json['file_id']}")
    assert dl_res.status_code == 200
    assert dl_res.data[:2] == b"PK"  # ZIP magic bytes


def test_convert_docx_to_pdf_endpoint(client):
    """Test uploading a Word DOCX file and converting to PDF."""
    doc = Document()
    doc.add_heading("Automated Test Document", level=0)
    doc.add_paragraph("This is paragraph text in docx.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_io.seek(0)

    data = {
        "conversion_type": "docx_to_pdf",
        "output_format": "pdf",
        "custom_filename": "converted_report",
        "files": (docx_io, "report.docx")
    }

    res = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert res.status_code == 200
    res_json = res.get_json()
    assert res_json["success"] is True
    assert res_json["filename"] == "converted_report.pdf"
    assert "file_id" in res_json

    dl_res = client.get(f"/download/{res_json['file_id']}")
    assert dl_res.status_code == 200
    assert dl_res.data.startswith(b"%PDF")


def test_convert_image_format_endpoint(client):
    """Test converting PNG image to WebP."""
    img_io = io.BytesIO()
    Image.new("RGB", (120, 120), color="green").save(img_io, format="PNG")
    img_io.seek(0)

    data = {
        "conversion_type": "image_converter",
        "output_format": "webp",
        "custom_filename": "converted_pic",
        "files": (img_io, "photo.png")
    }

    res = client.post("/api/convert", data=data, content_type="multipart/form-data")
    assert res.status_code == 200
    res_json = res.get_json()
    assert res_json["success"] is True
    assert res_json["filename"] == "converted_pic.webp"
