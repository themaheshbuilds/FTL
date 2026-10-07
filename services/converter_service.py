import os
import shutil
from typing import List, Optional, Tuple
from PIL import Image
from config import Config
from services.job_service import JobService
from services.pdf_service import PdfService
from services.docx_service import DocxService
from services.archive_service import ArchiveService
from utils.filenames import sanitize_filename
from utils.ffmpeg_helper import extract_audio, remux_video
from utils.logging import get_logger

logger = get_logger("linkforge.converter_service")


class ConverterService:
    """Orchestrates local file conversions (PDF, DOCX, Video, Audio, Images)."""

    @classmethod
    def convert_files(
        cls,
        conversion_type: str,
        input_files: List[str],
        output_format: Optional[str] = None,
        custom_filename: Optional[str] = None
    ) -> Tuple[bool, Optional[str], Optional[str], Optional[str], Optional[int], Optional[str]]:
        """Main dispatcher for file conversions.
        
        Returns:
            (success, job_id, filename, mime_type, filesize, error_message)
        """
        if not input_files:
            return False, None, None, None, None, "No input files provided for conversion."

        # Create isolated job
        job = JobService.create_job("local_upload", conversion_type)
        temp_dir = JobService.get_temp_dir(job.job_id)
        gen_dir = JobService.get_generated_dir(job.job_id)

        try:
            conv_type = (conversion_type or "").lower().strip()
            out_fmt = (output_format or "").lower().strip()

            # Base name
            first_fname = os.path.basename(input_files[0])
            raw_base = custom_filename.strip() if custom_filename and custom_filename.strip() else os.path.splitext(first_fname)[0]
            safe_base = sanitize_filename(raw_base, default_name="converted_file")

            # -------------------------------------------------------------
            # 1. PDF to Image (JPG / PNG or ZIP)
            # -------------------------------------------------------------
            if conv_type in ("pdf_to_images", "pdf_to_image", "pdf_to_jpg", "pdf_to_png"):
                import fitz
                pdf_path = input_files[0]
                img_ext = "png" if (out_fmt == "png" or "png" in conv_type) else "jpg"
                doc = fitz.open(pdf_path)
                num_pages = len(doc)
                if num_pages == 0:
                    return False, None, None, None, None, "PDF has no pages."

                extracted_paths = []
                pad = 2 if num_pages < 100 else 3
                for idx, page in enumerate(doc):
                    pix = page.get_pixmap(dpi=150)
                    img_name = f"{safe_base}_page_{idx + 1:0{pad}d}.{img_ext}"
                    img_dest = str(temp_dir / img_name)
                    pix.save(img_dest)
                    extracted_paths.append((img_dest, img_name))
                doc.close()

                if num_pages == 1:
                    # Single page: return image directly
                    single_path, single_name = extracted_paths[0]
                    final_path = str(gen_dir / single_name)
                    shutil.move(single_path, final_path)
                    fsize = os.path.getsize(final_path)
                    mime = "image/png" if img_ext == "png" else "image/jpeg"
                    JobService.update_job_success(job.job_id, final_path, single_name, mime, fsize)
                    return True, job.job_id, single_name, mime, fsize, None
                else:
                    # Multi page: package into ZIP
                    zip_name = f"{safe_base}_images.zip"
                    final_zip_path = str(gen_dir / zip_name)
                    ok = ArchiveService.create_zip(extracted_paths, final_zip_path)
                    if not ok or not os.path.exists(final_zip_path):
                        return False, None, None, None, None, "Failed to package PDF images into ZIP."
                    fsize = os.path.getsize(final_zip_path)
                    JobService.update_job_success(job.job_id, final_zip_path, zip_name, "application/zip", fsize)
                    return True, job.job_id, zip_name, "application/zip", fsize, None

            # -------------------------------------------------------------
            # 2. Images to PDF
            # -------------------------------------------------------------
            if conv_type in ("images_to_pdf", "image_to_pdf", "jpg_to_pdf", "png_to_pdf"):
                out_name = f"{safe_base}.pdf"
                out_path = str(gen_dir / out_name)
                ok = PdfService.images_to_pdf(input_files, out_path)
                if not ok or not os.path.exists(out_path):
                    return False, None, None, None, None, "Failed to compile images into PDF."
                fsize = os.path.getsize(out_path)
                JobService.update_job_success(job.job_id, out_path, out_name, "application/pdf", fsize)
                return True, job.job_id, out_name, "application/pdf", fsize, None

            # -------------------------------------------------------------
            # 3. DOCX to PDF
            # -------------------------------------------------------------
            if conv_type in ("docx_to_pdf", "doc_to_pdf", "word_to_pdf"):
                docx_path = input_files[0]
                out_name = f"{safe_base}.pdf"
                out_path = str(gen_dir / out_name)

                # Attempt 1: Word COM on Windows
                converted = False
                word = None
                wdoc = None
                try:
                    import win32com.client
                    import pythoncom
                    pythoncom.CoInitialize()
                    word = win32com.client.DispatchEx("Word.Application")
                    word.Visible = False
                    word.DisplayAlerts = 0
                    wdoc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
                    wdoc.SaveAs(os.path.abspath(out_path), FileFormat=17)  # wdFormatPDF
                    converted = os.path.exists(out_path) and os.path.getsize(out_path) > 0
                except Exception as com_err:
                    logger.warning(f"Word COM conversion failed, trying fallback: {com_err}")
                finally:
                    if wdoc is not None:
                        try:
                            wdoc.Close(False)
                        except Exception:
                            pass
                        del wdoc
                    if word is not None:
                        try:
                            word.Quit()
                        except Exception:
                            pass
                        del word
                    try:
                        import pythoncom
                        pythoncom.CoUninitialize()
                    except Exception:
                        pass

                # Attempt 2: Fallback ReportLab parser
                if not converted:
                    converted = cls._docx_to_pdf_fallback(docx_path, out_path)

                if not converted or not os.path.exists(out_path):
                    return False, None, None, None, None, "Could not convert DOCX to PDF."

                fsize = os.path.getsize(out_path)
                JobService.update_job_success(job.job_id, out_path, out_name, "application/pdf", fsize)
                return True, job.job_id, out_name, "application/pdf", fsize, None

            # -------------------------------------------------------------
            # 4. PDF to DOCX
            # -------------------------------------------------------------
            if conv_type in ("pdf_to_docx", "pdf_to_doc", "pdf_to_word"):
                pdf_path = input_files[0]
                out_name = f"{safe_base}.docx"
                out_path = str(gen_dir / out_name)

                ok = cls._pdf_to_docx(pdf_path, out_path)
                if not ok or not os.path.exists(out_path):
                    return False, None, None, None, None, "Failed to convert PDF to DOCX."

                fsize = os.path.getsize(out_path)
                mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                JobService.update_job_success(job.job_id, out_path, out_name, mime, fsize)
                return True, job.job_id, out_name, mime, fsize, None

            # -------------------------------------------------------------
            # 5. Video to Audio (MP3 / M4A)
            # -------------------------------------------------------------
            if conv_type in ("video_to_audio", "video_to_mp3", "mp4_to_mp3", "extract_audio"):
                video_path = input_files[0]
                target_audio_ext = "m4a" if (out_fmt == "m4a" or "m4a" in conv_type) else "mp3"
                out_name = f"{safe_base}.{target_audio_ext}"
                out_path = str(gen_dir / out_name)

                ok = extract_audio(video_path, out_path, target_format=target_audio_ext)
                if not ok or not os.path.exists(out_path):
                    return False, None, None, None, None, f"Failed to extract {target_audio_ext.upper()} audio."

                fsize = os.path.getsize(out_path)
                mime = "audio/mp4" if target_audio_ext == "m4a" else "audio/mpeg"
                JobService.update_job_success(job.job_id, out_path, out_name, mime, fsize)
                return True, job.job_id, out_name, mime, fsize, None

            # -------------------------------------------------------------
            # 6. Video Converter (MP4 / MKV / WebM)
            # -------------------------------------------------------------
            if conv_type in ("video_to_video", "video_converter", "mp4_to_mkv", "video_to_mkv", "video_to_webm"):
                video_path = input_files[0]
                target_container = "mkv" if (out_fmt == "mkv" or "mkv" in conv_type) else ("webm" if (out_fmt == "webm" or "webm" in conv_type) else "mp4")
                out_name = f"{safe_base}.{target_container}"
                out_path = str(gen_dir / out_name)

                ok = remux_video(video_path, out_path, target_format=target_container)
                if not ok or not os.path.exists(out_path):
                    return False, None, None, None, None, f"Failed to convert video to {target_container.upper()}."

                fsize = os.path.getsize(out_path)
                mime_map = {"mp4": "video/mp4", "mkv": "video/x-matroska", "webm": "video/webm"}
                mime = mime_map.get(target_container, "video/mp4")
                JobService.update_job_success(job.job_id, out_path, out_name, mime, fsize)
                return True, job.job_id, out_name, mime, fsize, None

            # -------------------------------------------------------------
            # 7. Image Format Converter (JPG / PNG / WebP)
            # -------------------------------------------------------------
            if conv_type in ("image_converter", "image_to_image", "jpg_to_png", "png_to_jpg", "image_to_webp"):
                img_path = input_files[0]
                target_ext = out_fmt if out_fmt in ("jpg", "jpeg", "png", "webp") else "png"
                if target_ext == "jpeg":
                    target_ext = "jpg"
                out_name = f"{safe_base}.{target_ext}"
                out_path = str(gen_dir / out_name)

                with Image.open(img_path) as im:
                    if target_ext in ("jpg", "jpeg") and im.mode in ("RGBA", "LA", "P"):
                        rgb = Image.new("RGB", im.size, (255, 255, 255))
                        rgb.paste(im, mask=im.split()[-1] if "A" in im.mode else None)
                        rgb.save(out_path, "JPEG", quality=95)
                    else:
                        save_fmt = "PNG" if target_ext == "png" else ("WEBP" if target_ext == "webp" else "JPEG")
                        im.save(out_path, save_fmt)

                if not os.path.exists(out_path):
                    return False, None, None, None, None, f"Failed to convert image to {target_ext.upper()}."

                fsize = os.path.getsize(out_path)
                mime = "image/png" if target_ext == "png" else ("image/webp" if target_ext == "webp" else "image/jpeg")
                JobService.update_job_success(job.job_id, out_path, out_name, mime, fsize)
                return True, job.job_id, out_name, mime, fsize, None

            return False, None, None, None, None, f"Unsupported conversion type '{conversion_type}'."

        except Exception as e:
            logger.error(f"Error executing conversion {conversion_type}: {e}", exc_info=True)
            return False, None, None, None, None, f"Conversion failed: {str(e)}"

    @classmethod
    def _pdf_to_docx(cls, pdf_path: str, output_docx_path: str) -> bool:
        """Convert PDF pages into formatted DOCX."""
        try:
            import fitz
            import docx
            from docx.shared import Inches, Pt
            import io

            doc_fitz = fitz.open(pdf_path)
            doc_word = docx.Document()

            for i, page in enumerate(doc_fitz):
                text = page.get_text()
                if text.strip():
                    for line in text.split("\n"):
                        clean_line = line.strip()
                        if clean_line:
                            p = doc_word.add_paragraph(clean_line)
                            p.paragraph_format.space_after = Pt(4)
                else:
                    # Scanned or image page
                    pix = page.get_pixmap(dpi=150)
                    img_bytes = pix.tobytes("png")
                    doc_word.add_picture(io.BytesIO(img_bytes), width=Inches(6.5))

                if i < len(doc_fitz) - 1:
                    doc_word.add_page_break()

            doc_fitz.close()
            doc_word.save(output_docx_path)
            return os.path.exists(output_docx_path) and os.path.getsize(output_docx_path) > 0
        except Exception as e:
            logger.warning(f"Error in _pdf_to_docx: {e}")
            return False

    @classmethod
    def _docx_to_pdf_fallback(cls, docx_path: str, output_pdf_path: str) -> bool:
        """Fallback DOCX to PDF converter using python-docx and ReportLab."""
        try:
            import docx
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors

            doc = docx.Document(docx_path)
            pdf = SimpleDocTemplate(
                output_pdf_path,
                pagesize=letter,
                rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
            )

            styles = getSampleStyleSheet()
            normal_style = styles["Normal"]
            elements = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    elements.append(Paragraph(safe_text, normal_style))
                    elements.append(Spacer(1, 8))

            if not elements:
                elements.append(Paragraph("Document", normal_style))

            pdf.build(elements)
            return os.path.exists(output_pdf_path) and os.path.getsize(output_pdf_path) > 0
        except Exception as e:
            logger.warning(f"Error in _docx_to_pdf_fallback: {e}")
            return False
